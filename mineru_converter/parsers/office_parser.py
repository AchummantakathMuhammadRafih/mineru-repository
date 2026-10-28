"""Native document converters for DOCX, PPTX, XLSX, and HTML to Markdown."""

from pathlib import Path
from typing import Optional, Tuple
import re

from mineru_converter.utils.logging_utils import get_logger

logger = get_logger()


def convert_docx_to_markdown(docx_path: Path, output_dir: Path) -> Path:
    """Convert DOCX file to Markdown using mammoth or python-docx."""
    docx_path = Path(docx_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    out_md = output_dir / f"{docx_path.stem}.md"

    # Attempt with mammoth (best formatting & headings)
    try:
        import mammoth
        with open(docx_path, "rb") as docx_file:
            result = mammoth.convert_to_markdown(docx_file)
            md_text = result.value
            out_md.write_text(md_text, encoding="utf-8")
            logger.info(f"✓ [DOCX Converter] Converted {docx_path.name} to Markdown via Mammoth.")
            return out_md
    except ImportError:
        pass
    except Exception as m_err:
        logger.warning(f"Mammoth DOCX conversion failed: {m_err}. Falling back to python-docx...")

    # Fallback with python-docx
    import docx
    doc = docx.Document(str(docx_path))
    md_lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name.lower() if para.style else ""
        if "heading 1" in style_name:
            md_lines.append(f"# {text}\n")
        elif "heading 2" in style_name:
            md_lines.append(f"## {text}\n")
        elif "heading 3" in style_name:
            md_lines.append(f"### {text}\n")
        elif "heading 4" in style_name:
            md_lines.append(f"#### {text}\n")
        elif "heading" in style_name:
            md_lines.append(f"##### {text}\n")
        else:
            md_lines.append(f"{text}\n")

    # Tables
    for table in doc.tables:
        table_rows = []
        for r_idx, row in enumerate(table.rows):
            row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            table_rows.append("| " + " | ".join(row_cells) + " |")
            if r_idx == 0:
                table_rows.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
        if table_rows:
            md_lines.append("\n" + "\n".join(table_rows) + "\n")

    out_md.write_text("\n".join(md_lines), encoding="utf-8")
    logger.info(f"✓ [DOCX Converter] Converted {docx_path.name} to Markdown via python-docx.")
    return out_md


def convert_pptx_to_markdown(pptx_path: Path, output_dir: Path) -> Path:
    """Convert PPTX presentation slides to Markdown with structured headings per slide."""
    import pptx

    pptx_path = Path(pptx_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    out_md = output_dir / f"{pptx_path.stem}.md"

    prs = pptx.Presentation(str(pptx_path))
    md_blocks = []

    for s_idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        slide_texts = []

        if slide.shapes.title and slide.shapes.title.text.strip():
            slide_title = slide.shapes.title.text.strip()

        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    p_txt = paragraph.text.strip()
                    if p_txt:
                        level = paragraph.level or 0
                        indent = "  " * level
                        slide_texts.append(f"{indent}- {p_txt}")
            elif shape.has_table:
                table_rows = []
                for r_idx, row in enumerate(shape.table.rows):
                    row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    table_rows.append("| " + " | ".join(row_cells) + " |")
                    if r_idx == 0:
                        table_rows.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
                if table_rows:
                    slide_texts.append("\n" + "\n".join(table_rows))

        heading = f"## Slide {s_idx}: {slide_title}" if slide_title else f"## Slide {s_idx}"
        slide_content = "\n".join(slide_texts)
        md_blocks.append(f"{heading}\n\n{slide_content}".strip())

    out_md.write_text("\n\n---\n\n".join(md_blocks), encoding="utf-8")
    logger.info(f"✓ [PPTX Converter] Converted {pptx_path.name} ({len(prs.slides)} slides) to Markdown.")
    return out_md


def convert_xlsx_to_markdown(xlsx_path: Path, output_dir: Path) -> Path:
    """Convert XLSX sheets to Markdown tables."""
    import openpyxl

    xlsx_path = Path(xlsx_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    out_md = output_dir / f"{xlsx_path.stem}.md"

    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    md_blocks = [f"# {xlsx_path.stem}\n"]

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        md_blocks.append(f"## Sheet: {sheet_name}\n")
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            md_blocks.append("*Empty sheet*\n")
            continue

        table_lines = []
        # Find maximum columns with data
        non_empty_rows = [r for r in rows if any(c is not None for c in r)]
        if not non_empty_rows:
            md_blocks.append("*Empty sheet*\n")
            continue

        max_cols = max(len(r) for r in non_empty_rows)
        for r_idx, row in enumerate(non_empty_rows):
            cells = [str(c if c is not None else "").strip().replace("\n", " ").replace("|", "/") for c in row]
            while len(cells) < max_cols:
                cells.append("")
            table_lines.append("| " + " | ".join(cells) + " |")
            if r_idx == 0:
                table_lines.append("| " + " | ".join(["---"] * max_cols) + " |")

        md_blocks.append("\n".join(table_lines) + "\n")

    out_md.write_text("\n\n".join(md_blocks), encoding="utf-8")
    logger.info(f"✓ [XLSX Converter] Converted {xlsx_path.name} to Markdown.")
    return out_md


def convert_html_to_markdown(html_path: Path, output_dir: Path) -> Path:
    """Convert HTML document into clean Markdown using BeautifulSoup."""
    from bs4 import BeautifulSoup

    html_path = Path(html_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    out_md = output_dir / f"{html_path.stem}.md"

    raw_html = html_path.read_text(encoding="utf-8", errors="replace")
    soup = BeautifulSoup(raw_html, "html.parser")

    # Remove script, style tags
    for tag in soup(["script", "style", "meta", "noscript"]):
        tag.decompose()

    md_lines = []

    # Title
    title_tag = soup.find("title")
    if title_tag and title_tag.text.strip():
        md_lines.append(f"# {title_tag.text.strip()}\n")

    # Traverse elements
    for el in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "table", "li", "pre", "blockquote"]):
        txt = el.get_text().strip()
        if not txt:
            continue

        if el.name == "h1":
            md_lines.append(f"# {txt}\n")
        elif el.name == "h2":
            md_lines.append(f"## {txt}\n")
        elif el.name == "h3":
            md_lines.append(f"### {txt}\n")
        elif el.name == "h4":
            md_lines.append(f"#### {txt}\n")
        elif el.name == "h5":
            md_lines.append(f"##### {txt}\n")
        elif el.name == "h6":
            md_lines.append(f"###### {txt}\n")
        elif el.name == "p":
            md_lines.append(f"{txt}\n")
        elif el.name == "blockquote":
            md_lines.append(f"> {txt}\n")
        elif el.name == "pre":
            md_lines.append(f"```\n{txt}\n```\n")
        elif el.name == "li":
            md_lines.append(f"- {txt}")
        elif el.name == "table":
            table_rows = []
            for r_idx, tr in enumerate(el.find_all("tr")):
                cols = [td.get_text().strip().replace("\n", " ") for td in tr.find_all(["td", "th"])]
                if cols:
                    table_rows.append("| " + " | ".join(cols) + " |")
                    if r_idx == 0:
                        table_rows.append("| " + " | ".join(["---"] * len(cols)) + " |")
            if table_rows:
                md_lines.append("\n" + "\n".join(table_rows) + "\n")

    out_md.write_text("\n".join(md_lines), encoding="utf-8")
    logger.info(f"✓ [HTML Converter] Converted {html_path.name} to Markdown.")
    return out_md
