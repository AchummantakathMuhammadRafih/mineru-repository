"""Tests for Office & HTML document converters (DOCX, PPTX, XLSX, HTML)."""

from pathlib import Path
import pytest

from mineru_converter.parsers.office_parser import (
    convert_docx_to_markdown,
    convert_pptx_to_markdown,
    convert_xlsx_to_markdown,
    convert_html_to_markdown,
)
from mineru_converter.core.converter import MinerUConverter
from mineru_converter.config.settings import ConversionConfig


def test_convert_html_to_markdown(tmp_path):
    """Test converting an HTML file to Markdown."""
    html_file = tmp_path / "test.html"
    html_file.write_text(
        "<html><head><title>Test Doc</title></head><body>"
        "<h1>Main Section</h1><p>This is a paragraph.</p>"
        "<table><tr><th>Col 1</th><th>Col 2</th></tr><tr><td>A</td><td>B</td></tr></table>"
        "</body></html>",
        encoding="utf-8"
    )

    out_dir = tmp_path / "out"
    md_path = convert_html_to_markdown(html_file, out_dir)

    assert md_path.exists()
    content = md_path.read_text(encoding="utf-8")
    assert "# Test Doc" in content or "# Main Section" in content
    assert "This is a paragraph." in content
    assert "| Col 1 | Col 2 |" in content


def test_convert_docx_to_markdown(tmp_path):
    """Test converting a DOCX file to Markdown via python-docx / mammoth."""
    import docx
    docx_file = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_heading("Document Title", level=1)
    doc.add_paragraph("First body paragraph.")
    doc.save(str(docx_file))

    out_dir = tmp_path / "out_docx"
    md_path = convert_docx_to_markdown(docx_file, out_dir)

    assert md_path.exists()
    content = md_path.read_text(encoding="utf-8")
    assert "Document Title" in content
    assert "First body paragraph" in content


def test_convert_xlsx_to_markdown(tmp_path):
    """Test converting an Excel sheet to Markdown."""
    import openpyxl
    xlsx_file = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Name", "Score"])
    ws.append(["Alice", 95])
    ws.append(["Bob", 88])
    wb.save(str(xlsx_file))

    out_dir = tmp_path / "out_xlsx"
    md_path = convert_xlsx_to_markdown(xlsx_file, out_dir)

    assert md_path.exists()
    content = md_path.read_text(encoding="utf-8")
    assert "## Sheet: Summary" in content
    assert "| Name | Score |" in content
    assert "| Alice | 95 |" in content


def test_converter_pipeline_with_docx(tmp_path):
    """Test full MinerUConverter pipeline with DOCX, including hierarchy and manifest."""
    import docx
    docx_file = tmp_path / "annual_report.docx"
    doc = docx.Document()
    doc.add_heading("Annual Report 2026", level=1)
    doc.add_paragraph("Confidential information for internal use.")
    doc.save(str(docx_file))

    cfg = ConversionConfig(
        output_dir=tmp_path / "output",
        redact_pii=False,
    )
    converter = MinerUConverter(config=cfg)
    result = converter.convert(docx_file)

    assert result.success is True
    assert result.manifest.markdown_path.exists()
    assert result.manifest.clean_markdown_path.exists()
    clean_md = result.manifest.clean_markdown_path.read_text(encoding="utf-8")
    assert "# Annual Report 2026" in clean_md
